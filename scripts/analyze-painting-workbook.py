"""Read-only extraction of the supplied painting source for audit/code review."""
import hashlib
import json
from pathlib import Path
import openpyxl

source = Path('C:/Users/ibetu/Desktop/BOYA_Son Hali.xlsx')
destination = Path('test-results/painting/source-audit.json')
formulas = openpyxl.load_workbook(source, read_only=True, data_only=False)
values = openpyxl.load_workbook(source, read_only=True, data_only=True)
summary = formulas['Fiyat Özeti']
price_values = values['Fiyat Özeti']
analysis = values['Analiz Basitleştirilmiş']
records = []
for row in range(5, 26):
    name_ref = summary[f'A{row}'].value
    price_ref = summary[f'B{row}'].value
    if not isinstance(name_ref, str) or not isinstance(price_ref, str):
        raise ValueError(f'Missing formula row {row}')
    name_cell = name_ref.split('!')[-1]
    price_cell = price_ref.split('!')[-1]
    name = analysis[name_cell].value
    cached_name = price_values[f'A{row}'].value
    price = price_values[f'B{row}'].value
    cached_price = analysis[price_cell].value
    if name != cached_name or price != cached_price or not isinstance(price, (int, float)):
        raise ValueError(f'Summary/analysis contradiction in row {row}')
    last_row = analysis[price_cell].row
    first_row = analysis[name_cell].row
    ingredients = []
    for ingredient_row in range(first_row + 1, last_row):
        label = analysis[f'A{ingredient_row}'].value
        amount = analysis[f'F{ingredient_row}'].value
        if isinstance(label, str) and isinstance(amount, (int, float)):
            ingredients.append({
                'row': ingredient_row,
                'label': label.strip(),
                'unit': analysis[f'C{ingredient_row}'].value,
                'quantity': analysis[f'D{ingredient_row}'].value,
                'unitPrice': analysis[f'E{ingredient_row}'].value,
                'amount': amount,
            })
    records.append({
        'summaryRow': row,
        'nameCell': name_cell,
        'priceCell': price_cell,
        'description': name.strip(),
        'unitPriceM2': price,
        'priceFormula': formulas['Analiz Basitleştirilmiş'][price_cell].value,
        'ingredients': ingredients,
    })
report = {
    'source': str(source),
    'sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
    'sheets': [{'name': sheet.title, 'rows': sheet.max_row, 'cols': sheet.max_column} for sheet in formulas],
    'positions': records,
    'questionTree': [
        {'row': row, 'text': values['Soru Ağacı'][f'A{row}'].value}
        for row in range(1, 46) if values['Soru Ağacı'][f'A{row}'].value is not None
    ],
}
destination.parent.mkdir(parents=True, exist_ok=True)
destination.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(f'positions={len(records)} sheets={len(report["sheets"])} output={destination}')
